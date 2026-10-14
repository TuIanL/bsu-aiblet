# -*- coding: utf-8 -*-
"""由 form_content.json 生成《安行护腰 · 第二点位执行记录表》Word 文档。

同一套内容渲染两份：
  · 空白模板  —— 供第二点位执行时填写
  · 填写示例  —— 取首个点位实际执行情况，供格式与口径参照
"""
import copy
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from inline import add_inline

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = json.load(open(os.path.join(HERE, "form_content.json"), encoding="utf-8"))

BLUE = RGBColor(0x0E, 0x6F, 0xA8)
DARK = RGBColor(0x1A, 0x2E, 0x3C)
GREY = RGBColor(0x60, 0x66, 0x70)
RED = RGBColor(0xC0, 0x39, 0x39)
ORANGE = RGBColor(0xC9, 0x6A, 0x1C)
GREEN = RGBColor(0x1E, 0x7A, 0x52)


# --------------------------------------------------------------- 排版工具
def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(4)
    for lvl, (sz, col) in enumerate([(18, BLUE), (14, BLUE), (12, DARK), (10.5, GREY)], start=1):
        s = doc.styles[f"Heading {lvl}"]
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = col
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        s.paragraph_format.space_before = Pt(14 if lvl <= 2 else 9)
        s.paragraph_format.space_after = Pt(6 if lvl <= 2 else 4)
        s.paragraph_format.line_spacing = 1.3
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.0)
    return doc


def P(doc, text="", size=10.5, bold=False, color=None, align=None,
      before=0, after=4, spacing=1.5, cn="宋体"):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    add_inline(p, text, size=size, bold=bold, color=color, cn=cn)
    return p


def H(doc, text, lvl=1):
    p = doc.add_heading(text, level=lvl)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return p


def shade(cell, hexc):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexc)
    cell._tc.get_or_add_tcPr().append(el)


def borders(t, color="BFBFBF", sz=4):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    t._tbl.tblPr.append(b)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def ctext(cell, txt, font, bold=False, center=False, color=None, cn="宋体"):
    cell.text = ""
    for i, seg in enumerate(str(txt).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.2
        add_inline(p, seg, size=font, bold=bold, color=color, cn=cn)


def TBL(doc, headers, rows, widths=None, font=8.5, center=(), caption=None,
        head_fill="DCEAF3", zebra="F7FAFC"):
    if caption:
        P(doc, caption, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
          before=7, after=3, cn="黑体")
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        ctext(hdr.cells[i], str(h), font, bold=True, center=True, color=DARK, cn="黑体")
        shade(hdr.cells[i], head_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            ctext(cells[i], v, font, center=(i in center))
            if zebra and ri % 2 == 1:
                shade(cells[i], zebra)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def BOX(doc, lines, fill="F2F7FA", border="B8D4E5", title=None, tcolor=BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, fill)
    borders(t, color=border, sz=6)
    c.text = ""
    seq = ([title] if title else []) + list(lines)
    for i, ln in enumerate(seq):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.4
        add_inline(p, ln, size=9.5, bold=(i == 0 and bool(title)),
                   color=tcolor if (i > 0 or not title) else DARK, cn="宋体")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def pagebreak(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------- 渲染一份
def render(data, mode):
    """mode: 'template' 空白模板 | 'example' 填写示例"""
    filled = (mode == "example")
    doc = new_doc()
    M = data["meta"]

    # ---------------- 封面区
    P(doc)
    P(doc, M["title"], size=24, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER,
      after=3, cn="黑体", spacing=1.1)
    P(doc, M["subtitle"], size=17, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER,
      after=4, cn="黑体", spacing=1.2)
    P(doc, M["en"], size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    P(doc, M["purpose"], size=11, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    P(doc, M["companion"], size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    if filled:
        BOX(doc, ["本文档为**填写示例**，内容取自首个点位（颐清园养老社区服务驿站）的实际执行情况，"
                  "仅用于说明填写方式与口径。第二点位执行时请使用《空白模板》版本，"
                  "按实际情况重新填写，**不得直接沿用本文档数据**。"],
            fill="FFF6E8", border="F0C896", title="示例文档说明", tcolor=ORANGE)
    else:
        BOX(doc, ["本表为**空白模板**。请用黑色签字笔填写；不适用的项划斜线并签字确认。"
                  "带 ** 标记的为关键项，漏填将使本表失去证明效力。"],
            fill="EDF7F1", border="A8D8BF", title="模板使用提示", tcolor=GREEN)

    # ---------------- 基本信息
    H(doc, "表单基本信息", 2)
    flds = M["fields"]
    ex = data.get("_example_fields") if filled else None
    if filled:
        ex = data["part1"]["example"]["fields"]
    rows = []
    for i, fld in enumerate(flds):
        if filled:
            ex_row = ex[i]
            fld = [fld[0], ex_row[1], fld[2], ex_row[3]]
        rows.append(list(fld))
    TBL(doc, ["项目", "内容", "项目", "内容"], rows,
        widths=[2.6, 5.4, 2.6, 5.4], font=9, caption=None)
    P(doc, after=4)

    DECL = M["declaration"]
    BOX(doc, DECL["lines"], title=DECL["title"], fill="F2F7FA", border="B8D4E5")
    pagebreak(doc)

    # ---------------- 第一部分
    P1 = data["part1"]
    H(doc, P1["title"], 1)
    P(doc, P1["subtitle"], size=9.5, color=GREY, after=8)
    ex1 = P1["example"] if filled else None
    for si, sec in enumerate(P1["sections"]):
        H(doc, sec["h"], 2)
        rows = sec["rows"]
        if filled:
            key = {0: "s11", 1: "s12", 2: "s13", 3: "s14"}[si]
            rows = ex1[key]
        TBL(doc, sec["headers"], rows, widths=sec.get("widths"),
            center=tuple(range(len(sec["headers"]))) if filled and si in (1, 2, 3) else (),
            head_fill="DCEAF3" if filled else "E8E8E8")
        P(doc, after=4)
    BOX(doc, [P1["note"]], title="本部分的证明作用", fill="F2F7FA", border="B8D4E5")
    pagebreak(doc)

    # ---------------- 第二部分
    P2 = data["part2"]
    H(doc, P2["title"], 1)
    P(doc, P2["subtitle"], size=9.5, color=GREY, after=6)
    P(doc, P2["intro"], size=10)
    ex2 = P2["example"] if filled else None
    gkeys = ["g1", "g2", "g3", "g4", "g5"]
    for gi, grp in enumerate(P2["groups"]):
        H(doc, grp["h"], 2)
        rows = ex2[gkeys[gi]] if filled else grp["rows"]
        TBL(doc, grp["headers"], rows, widths=grp.get("widths"),
            center=tuple(grp.get("center", [])),
            head_fill="DCEAF3" if filled else "E8E8E8")
        P(doc, after=4)
    H(doc, P2["summary"]["h"], 2)
    S = P2["summary"]
    TBL(doc, S["headers"], ex2["sum"] if filled else S["rows"],
        widths=S.get("widths"), center=tuple(S.get("center", [])),
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, after=4)
    BOX(doc, [P2["note"]], title="本部分的证明作用", fill="F2F7FA", border="B8D4E5")
    pagebreak(doc)

    # ---------------- 第三部分
    P3 = data["part3"]
    H(doc, P3["title"], 1)
    P(doc, P3["subtitle"], size=9.5, color=GREY, after=8)
    ex3 = P3["example"] if filled else None
    H(doc, P3["sections"][0]["h"], 2)
    TBL(doc, P3["sections"][0]["headers"],
        ex3["s1"] if filled else P3["sections"][0]["rows"],
        widths=P3["sections"][0].get("widths"), center=(0, 2, 3),
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, after=4)
    H(doc, P3["sections"][1]["h"], 2)
    TBL(doc, P3["sections"][1]["headers"],
        ex3["s2"] if filled else P3["sections"][1]["rows"],
        widths=P3["sections"][1].get("widths"), center=(2, 3, 4),
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, after=4)
    H(doc, P3["sections"][2]["h"], 2)
    TBL(doc, P3["sections"][2]["headers"],
        ex3["s3"] if filled else P3["sections"][2]["rows"],
        widths=P3["sections"][2].get("widths"), center=(3,),
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, after=4)
    H(doc, P3["sections"][3]["h"], 2)
    TBL(doc, P3["sections"][3]["headers"],
        ex3["s4"] if filled else P3["sections"][3]["rows"],
        widths=P3["sections"][3].get("widths"),
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, after=4)
    BOX(doc, [P3["note"]], title="本部分的证明作用", fill="F2F7FA", border="B8D4E5")
    P(doc, after=6)

    # ---------------- 签署
    F = data["footer"]
    H(doc, "签署", 2)
    sig_rows = [[r, "", "", ""] for r in F["sign"]]
    if filled:
        sig_rows = [[F["sign"][0], "×××", "（签字）", "2026年10月20日"],
                    [F["sign"][1], "×××", "（签字）", "2026年10月20日"],
                    [F["sign"][2], "×××", "（签字）", "2026年10月21日"]]
    TBL(doc, ["角色", "姓名", "签字", "日期"], sig_rows,
        widths=[4.6, 3.4, 4.0, 3.0], font=9,
        head_fill="DCEAF3" if filled else "E8E8E8")
    P(doc, F["sign_note"], size=9.5, color=GREY, after=8)
    P(doc, F["note"], size=8.5, color=GREY, spacing=1.4)

    suffix = "（填写示例）" if filled else "（空白模板）"
    out = os.path.join(HERE, f"安行护腰_第二点位执行记录表_{suffix}.docx")
    doc.save(out)
    return out


paths = [render(SRC, "template"), render(SRC, "example")]
for p in paths:
    print("saved:", os.path.basename(p), os.path.getsize(p), "bytes")
