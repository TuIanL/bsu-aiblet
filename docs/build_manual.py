# -*- coding: utf-8 -*-
"""由 manual_content.json 生成《安行护腰 · 社区落地标准化作业手册》Word 文档。

内容与排版分离：内容全在 JSON 里，本脚本只负责排版。
好处是内容改动不需要碰 Python 字符串，规避中文引号转义问题。
"""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from inline import add_inline

HERE = os.path.dirname(os.path.abspath(__file__))
D = json.load(open(os.path.join(HERE, "manual_content.json"), encoding="utf-8"))

BLUE = RGBColor(0x0E, 0x6F, 0xA8)
DARK = RGBColor(0x1A, 0x2E, 0x3C)
GREY = RGBColor(0x60, 0x66, 0x70)
RED = RGBColor(0xC0, 0x39, 0x39)

TONE = {
    "F2F7FA|B8D4E5": RGBColor(0x0E, 0x6F, 0xA8),   # 蓝：口径声明
    "EDF7F1|A8D8BF": RGBColor(0x1E, 0x7A, 0x52),   # 绿：结论与定额
    "FFF6E8|F0C896": RGBColor(0xC9, 0x6A, 0x1C),   # 橙：提醒与风险
    "FDEEEE|EFC0C0": RGBColor(0xC0, 0x39, 0x39),   # 红：红线警示
    "F6F6F6|CCCCCC": RGBColor(0x60, 0x66, 0x70),   # 灰：一般说明
}

doc = Document()

# ------------------------------------------------------------------ 基础样式
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(4)

for lvl, (sz, col) in enumerate([(18, BLUE), (14, BLUE), (12, DARK), (10.5, GREY)], start=1):
    s = doc.styles[f"Heading {lvl}"]
    s.font.name = "Times New Roman"
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = col
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    s.paragraph_format.space_before = Pt(14 if lvl <= 2 else 9)
    s.paragraph_format.space_after = Pt(6 if lvl <= 2 else 4)
    s.paragraph_format.line_spacing = 1.3

sec = doc.sections[0]
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.2)


def P(text="", size=10.5, bold=False, color=None, align=None, before=0, after=4,
      spacing=1.5, cn="宋体"):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    add_inline(p, text, size=size, bold=bold, color=color, cn=cn)
    return p


def H(text, lvl=1):
    p = doc.add_heading(text, level=lvl)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return p


def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexc)
    tcPr.append(el)


def borders(t, color="BFBFBF", sz=4):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    t._tbl.tblPr.append(b)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def cell_text(cell, txt, font, bold=False, center=False, color=None, cn="宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.2
    add_inline(p, str(txt), size=font, bold=bold, color=color, cn=cn)


def table(headers, rows, widths=None, font=9, center=(), caption=None):
    if caption:
        P(caption, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
          before=8, after=3, cn="黑体")
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        cell_text(hdr.cells[i], str(h), font, bold=True, center=True, color=DARK, cn="黑体")
        shade(hdr.cells[i], "DCEAF3")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cell_text(cells[i], str(v), font, center=(i in center))
            if ri % 2 == 1:
                shade(cells[i], "F7FAFC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def box(lines, fill="F2F7FA", border="B8D4E5", title=None):
    tcolor = TONE.get(f"{fill}|{border}", DARK)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, fill)
    borders(t, color=border, sz=6)
    c.text = ""
    seq = list(lines)
    if title:
        seq = [title] + seq
    for i, ln in enumerate(seq):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.4
        add_inline(p, ln, size=9.5, bold=(i == 0 and bool(title)),
                   color=tcolor if (i > 0 or not title) else DARK, cn="宋体")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullets(items):
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.0)
        pf.first_line_indent = Cm(-0.5)
        pf.line_spacing = 1.45
        pf.space_after = Pt(2)
        add_inline(p, "● " + it, size=10.5)


def pagebreak():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================== 封面
M = D["meta"]
P()
P()
P(M["title"], size=26, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER,
  after=4, cn="黑体", spacing=1.1)
P(M["subtitle"], size=19, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER,
  after=6, cn="黑体", spacing=1.2)
P(M["en"], size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
P(M["tagline"], size=12, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
P(M["claim"], size=10.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)

t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in M["info"]:
    cells = t.add_row().cells
    for i, (txt, bold) in enumerate([(k, True), (v, False)]):
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(txt)
        r.bold = bold
        r.font.size = Pt(10.5)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if bold:
            r.font.color.rgb = DARK
    cells[0].width = Cm(3.4)
    cells[1].width = Cm(11.6)

P()
P()
DECL = M["declaration"]
box(DECL["lines"], fill="F2F7FA", border="B8D4E5", title=DECL["title"])
pagebreak()

# ================================================================== 目录
H("目录", 1)
P()
for no, name in D["toc"]:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.7
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{no}　{name}")
    r.font.size = Pt(10.5)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if "." not in no:
        r.bold = True
        r.font.color.rgb = DARK
P()
P(D["toc_note"], size=9.5, color=GREY)
pagebreak()

# ================================================================== 正文
for b in D["blocks"]:
    t_ = b["type"]
    if t_ == "h1":
        H(b["text"], 1)
    elif t_ == "h2":
        H(b["text"], 2)
    elif t_ == "p":
        bold = b.get("bold", False)
        txt = b["text"]
        if txt.startswith("**") and txt.endswith("**"):
            txt, bold = txt[2:-2], True
        P(txt, size=b.get("size", 10.5), bold=bold,
          color=RGBColor(0x60, 0x66, 0x70) if b.get("color") else None)
    elif t_ == "bullets":
        bullets(b["items"])
    elif t_ == "box":
        box(b["lines"], fill=b.get("fill", "F2F7FA"), border=b.get("border", "B8D4E5"),
            title=b.get("title"))
    elif t_ == "table":
        table(b["headers"], b["rows"], widths=b.get("widths"),
              font=b.get("font", 9), center=tuple(b.get("center", [])), caption=b.get("caption"))
        P(after=4)
    elif t_ == "pagebreak":
        pagebreak()

# ================================================================== 尾页
F = D["footer"]
P()
P(F["end"], size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
P(F["note"], size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

out = os.path.join(HERE, "安行护腰_社区落地标准化作业手册_V1.0.docx")
doc.save(out)
print("saved:", out)
print("size:", os.path.getsize(out), "bytes")
